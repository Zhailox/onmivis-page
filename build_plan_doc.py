import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Segoe UI'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = RGBColor(30, 41, 59) # Slate 800

    # Color Palette Constants
    COLOR_NAVY = RGBColor(15, 32, 66)       # #0F2042
    COLOR_CYAN = RGBColor(0, 168, 181)      # #00A8B5
    COLOR_PURPLE = RGBColor(131, 95, 158)   # #835F9E
    COLOR_DARK = RGBColor(2, 7, 19)         # #020713

    # ==========================
    # PORTADA / HEADER DEL DOCUMENTO
    # ==========================
    p_pre = doc.add_paragraph()
    p_pre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_pre = p_pre.add_run("OMNIVIS LABS // PLAN MAESTRO DE DESARROLLO")
    run_pre.font.name = 'Segoe UI'
    run_pre.font.size = Pt(10)
    run_pre.font.bold = True
    run_pre.font.color.rgb = COLOR_CYAN

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(8)
    run_title = p_title.add_run("LANDING PAGE OFICIAL:\nOMNIVIS & ESTACIÓN VERSALLES")
    run_title.font.name = 'Segoe UI'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_NAVY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(24)
    run_sub = p_sub.add_run("Arquitectura de Software, Sistema de Diseño Maximalista Espacial,\nEstrategia Promocional y Despliegue en Vercel")
    run_sub.font.name = 'Segoe UI'
    run_sub.font.size = Pt(11.5)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    # Box Metadata
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_data = [
        [("Marca & Estudio:", " Omnivis (Intelligence qui transforme votre environnement)"), 
         ("Producto Principal:", " Estación Versalles (Simulador Espacial & CLI)")],
        [("Stack Tecnológico:", " React 19 + TypeScript + Tailwind v4 + Motion + Vercel"), 
         ("Inspiración Estética:", " Byotōne (landing.love) / Maximalismo Espacial Sci-Fi")]
    ]

    for r_idx, row in enumerate(meta_data):
        for c_idx, (label, val) in enumerate(row):
            cell = meta_table.cell(r_idx, c_idx)
            set_cell_background(cell, "F1F5F9")
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r1 = p.add_run(label)
            r1.bold = True
            r1.font.size = Pt(9.5)
            r1.font.color.rgb = COLOR_NAVY
            r2 = p.add_run(val)
            r2.font.size = Pt(9.5)
            r2.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # ==========================
    # SECCIÓN 1: INTRODUCCIÓN Y OBJETIVOS
    # ==========================
    h1 = doc.add_heading("1. Resumen Ejecutivo y Objetivos Estratégicos", level=1)
    h1.style.font.color.rgb = COLOR_NAVY
    h1.style.font.name = 'Segoe UI'

    doc.add_paragraph(
        "El presente proyecto comprende la concepción, diseño e implementación de la Landing Page Oficial de la marca Omnivis "
        "y su videojuego insignia, 'Estación Versalles'. La landing page cumple un rol promocional y de conversión clave: "
        "comunica la identidad visionaria de Omnivis y posiciona a 'Estación Versalles' como un simulador de vanguardia que combina "
        "ingeniería aeroespacial, resolución de crisis bajo presión y programación determinista en Python y Java."
    )

    doc.add_paragraph(
        "La experiencia web está concebida bajo una filosofía maximalista inspirada en los estándares de Landing.love (específicamente "
        "la estética fluida e inmersiva de 'Byotōne'), incorporando micro-transiciones reactivas, efectos de partículas orbitales, "
        "audio ambiental interactivo ('Evolution of the Machine'), un reproductor táctico y un enlace directo de juego instantáneo "
        "al despliegue en Vercel."
    )

    # Bullet points
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.left_indent = Inches(0.2)
    p_obj.add_run("• Proporción de Contenido Equilibrada: ").bold = True
    p_obj.add_run("Sección inicial enfocada en la filosofía y visión de Omnivis, seguida de un despliegue de alto impacto promocional centrado en Estación Versalles.\n")
    p_obj.add_run("• Micro-Interacciones & Maximalismo Futurista: ").bold = True
    p_obj.add_run("Efectos de resplandor cuántico, bordes con degradados animados, tarjetas holográficas con inclinación 3D (tilt) y telemetría en tiempo real.\n")
    p_obj.add_run("• Banda Sonora Inmersiva Integrada: ").bold = True
    p_obj.add_run("Reproducción continua en bucle del tema oficial 'Evolution of the Machine.mp3' a volumen atenuado con control de usuario (play/pause, volumen y visualizador de frecuencias).\n")
    p_obj.add_run("• Despliegue en Vercel de Cero Fricción: ").bold = True
    p_obj.add_run("Arquitectura frontend SPA pura (React + Vite + Tailwind + Motion) 100% compatible con los límites y optimizaciones del CDN global de Vercel.")

    # ==========================
    # SECCIÓN 2: IDENTIDAD VISUAL Y PALETA DE COLORES
    # ==========================
    h2 = doc.add_heading("2. Sistema de Diseño e Identidad Cromática", level=1)
    h2.style.font.color.rgb = COLOR_NAVY

    doc.add_paragraph(
        "La paleta de colores y el lenguaje visual se han extraído y calibrado directamente a partir del logo oficial de Omnivis "
        "(el ojo cósmico en vórtice con gradiente bicromático) y los activos artísticos del juego ubicados en la carpeta del proyecto:"
    )

    # Color Table
    color_table = doc.add_table(rows=6, cols=3)
    color_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    color_table.autofit = False

    headers = ["Tono / Rol", "Código HEX", "Aplicación en la Interfaz"]
    for i, h in enumerate(headers):
        cell = color_table.cell(0, i)
        set_cell_background(cell, "0F2042")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)

    color_rows = [
        ("Void Deep Space (Fondo Base)", "#020713", "Fondo primordial de la página, contraste absoluto para elementos luminiscentes."),
        ("Omnivis Navy Blue (Marca)", "#15355A / #0F2042", "Contenedores estructurales, gradientes del hero y headers con desenfoque de cristal."),
        ("Electric Quantum Cyan (Acento)", "#00F5FF / #00C2CB", "Botones CTA primarios, bordes de neón activos, estados hover y láseres de escaneo."),
        ("Iridescent Nebula Violet (Marca)", "#835F9E / #A378C6", "Acentos secundarios tomados de la espiral superior del logo, gradientes de texto e insignias."),
        ("Stellar White & Chrome (Tipografía)", "#FFFFFF / #E2E8F0", "Títulos monumentales, textos de telemetría y detalles metálicos HUD.")
    ]

    for r_idx, (name, hex_code, usage) in enumerate(color_rows, start=1):
        c0 = color_table.cell(r_idx, 0)
        c1 = color_table.cell(r_idx, 1)
        c2 = color_table.cell(r_idx, 2)
        
        bg_col = "F8FAFC" if r_idx % 2 == 1 else "EDF2F7"
        set_cell_background(c0, bg_col)
        set_cell_background(c1, bg_col)
        set_cell_background(c2, bg_col)
        set_cell_margins(c0, 80, 80, 100, 100)
        set_cell_margins(c1, 80, 80, 100, 100)
        set_cell_margins(c2, 80, 80, 100, 100)

        p0 = c0.paragraphs[0]; r0 = p0.add_run(name); r0.bold = True; r0.font.size = Pt(9)
        p1 = c1.paragraphs[0]; r1 = p1.add_run(hex_code); r1.font.bold = True; r1.font.size = Pt(9); r1.font.color.rgb = COLOR_CYAN
        p2 = c2.paragraphs[0]; r2 = p2.add_run(usage); r2.font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Tipografía
    doc.add_paragraph(
        "Sistema Tipográfico Recomendado:\n"
        "• Orbitron: Para títulos monumentales, jerarquías de mando y encabezados de sector.\n"
        "• Rajdhani / Space Grotesk: Para navegación, insignias de estado y subtítulos tácticos.\n"
        "• JetBrains Mono / Fira Code: Para telemetría de vuelo, fragmentos de código en Python/Java y comandos de consola."
    )

    # ==========================
    # SECCIÓN 3: ESTRUCTURA DE LA PÁGINA Y EXPERIENCIA DE USUARIO
    # ==========================
    h3 = doc.add_heading("3. Estructura y Recorrido de la Landing Page", level=1)
    h3.style.font.color.rgb = COLOR_NAVY

    doc.add_paragraph(
        "La navegación está organizada en un flujo continuo con transiciones de scroll que guían al usuario desde la revelación "
        "de la marca hasta la adrenalina táctica del juego:"
    )

    sections_flow = [
        ("A. Header HUD Flotante con Movimiento", 
         "Barra superior de navegación fija con backdrop-filter glassmorphism, indicador de pulso de telemetría, "
         "logo de Omnivis con micro-animación en hover, selector de secciones rápidas, widget de audio 'Evolution of the Machine' "
         "(con visualizador en tiempo real y botón de mute) y botón de acceso rápido 'DESPLEGAR SIMULADOR'."),

        ("B. Bloque 1: Manifiesto Omnivis ('Intelligence qui transforme votre environnement')", 
         "Sección introductoria elegante que presenta a Omnivis como el laboratorio creador. Expone los principios fundacionales "
         "documentados en la estrategia comercial: soberanía tecnológica, diseño determinista sin modelos inflados, "
         "pedagogía orientada a la ingeniería dura y resiliencia offline."),

        ("C. Bloque 2: Hero Section Monumental — Estación Versalles", 
         "Impacto visual máximo: tipografía cinética espacial, selector bilingüe conmemorativo (ES 🇻🇪 / FR 🇫🇷), "
         "vitrina interactiva 3D con inclinación de la Estación Versalles ('Mapa General Estación Versalles.webp'), "
         "insignias de estado ('EN LÍNEA // SECTOR ORBITAL ALPHA') y el Call To Action estelar con resplandor neón: 'JUGAR EN VERCEL'."),

        ("D. Bloque 3: Pilares de Jugabilidad y Simulación Aeroespacial", 
         "Tarjetas maximalistas con micro-interacciones que exponen las mecánicas nucleares: "
         "Consola UNIX determinista con soporte real para scripts de Python y Java, cartografía hiperespacial (Versalles, Hyperion, Titán-IV, Helios), "
         "y el vertiginoso 'Modo Crisis' (fusión de plasma, desestabilización magnética, inyección de malware y decaimiento orbital)."),

        ("E. Bloque 4: Galería Táctica de Misiones e Instalaciones", 
         "Explorador interactivo de sectores espaciales con los artes oficiales del proyecto: "
         "Puente de Navegación, Vista Laboratorio de Datos, Vista Núcleo Reactor, Estación de Drones, Matriz de Escudos y los Acorazados Hyperion y Titán."),

        ("F. Bloque 5: Mini Terminal Interactivo 'Void CLI'", 
         "Un componente de micro-experiencia donde el visitante puede ingresar comandos reales en la landing (`man`, `diagnostico`, `play`, `omnivis`, `crew`) "
         "recibiendo respuestas sonoras y visuales inmediatas antes de ingresar a la versión completa."),

        ("G. Bloque 6: Audio Hub Oficial ('Evolution of the Machine')", 
         "Sección dedicada a la banda sonora del juego con tarjeta holográfica del tema, reproductor gráfico de frecuencias, "
         "información de la pista a 320 kbps y selector de pistas adicionales del soundtrack."),

        ("H. Bloque 7: Lanzadera Final (Call To Action Vercel)", 
         "Módulo de conversión final con animación de hiperespacio convocando a pilotos, ingenieros y programadores a asumir el mando del simulador."),

        ("I. Footer Maximalista con Enlaces Sociales y Telemetría", 
         "Pie de página con logotipo de Omnivis, enlaces a redes sociales (GitHub, Vercel, X/Twitter, Discord, LinkedIn), "
         "licencia de código abierto MIT, coordenadas orbitales y créditos de autoría.")
    ]

    for title, desc in sections_flow:
        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.space_before = Pt(6)
        p_sec.paragraph_format.space_after = Pt(4)
        run_sec_title = p_sec.add_run(title)
        run_sec_title.bold = True
        run_sec_title.font.size = Pt(11)
        run_sec_title.font.color.rgb = COLOR_PURPLE

        p_desc = doc.add_paragraph(desc)
        p_desc.paragraph_format.left_indent = Inches(0.2)
        p_desc.paragraph_format.space_after = Pt(8)

    # ==========================
    # SECCIÓN 4: ARQUITECTURA TÉCNICA Y VERCEL
    # ==========================
    h4 = doc.add_heading("4. Arquitectura de Implementación y Compatibilidad Vercel", level=1)
    h4.style.font.color.rgb = COLOR_NAVY

    doc.add_paragraph(
        "Para garantizar tiempos de carga sub-segundo, animaciones a 60 FPS y despliegue sin fallos en Vercel, se ha seleccionado "
        "la siguiente arquitectura modular:"
    )

    tech_table = doc.add_table(rows=6, cols=3)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tech_table.autofit = False

    t_headers = ["Tecnología / Librería", "Versión", "Rol en el Proyecto"]
    for i, h in enumerate(t_headers):
        cell = tech_table.cell(0, i)
        set_cell_background(cell, "0F2042")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)

    tech_rows = [
        ("React 19 + TypeScript", "v19.x / v5.x", "Ecosistema reactivo fuertemente tipado para estabilidad de componentes y modularidad."),
        ("Vite 6 / 8", "Última estable", "Bundler ultraveloz, hot reload instantáneo y optimización de activos estáticos para Vercel."),
        ("Tailwind CSS v4", "v4.x", "Motor de utilidades CSS de alto rendimiento con variables de tema personalizadas para la paleta Omnivis."),
        ("Motion (Framer Motion)", "v12.x", "Gestión declarativa de micro-transiciones, scroll-linked animations y orquestación de entrada."),
        ("Web Audio API / HTML5 Audio", "Nativo", "Reproducción en bucle con fade-in automático, analizador de nodos para el visualizador y toggle accesible.")
    ]

    for r_idx, (tech, ver, rol) in enumerate(tech_rows, start=1):
        c0 = tech_table.cell(r_idx, 0); c1 = tech_table.cell(r_idx, 1); c2 = tech_table.cell(r_idx, 2)
        bg = "F8FAFC" if r_idx % 2 == 1 else "EDF2F7"
        set_cell_background(c0, bg); set_cell_background(c1, bg); set_cell_background(c2, bg)
        set_cell_margins(c0, 80, 80, 100, 100); set_cell_margins(c1, 80, 80, 100, 100); set_cell_margins(c2, 80, 80, 100, 100)
        p0 = c0.paragraphs[0]; r0 = p0.add_run(tech); r0.bold = True; r0.font.size = Pt(9)
        p1 = c1.paragraphs[0]; r1 = p1.add_run(ver); r1.font.size = Pt(9); r1.font.color.rgb = COLOR_PURPLE
        p2 = c2.paragraphs[0]; r2 = p2.add_run(rol); r2.font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ==========================
    # SECCIÓN 5: PLAN DE EJECUCIÓN PASO A PASO
    # ==========================
    h5 = doc.add_heading("5. Fases de Ejecución del Proyecto", level=1)
    h5.style.font.color.rgb = COLOR_NAVY

    phases = [
        ("Fase 1: Andamiaje y Transferencia de Activos",
         "Creación del proyecto en Vite + React + TypeScript + Tailwind CSS. "
         "Transferencia de todos los recursos gráficos (logo.webp, ilustraciones de misiones, naves) y el audio 'Evolution of the Machine.mp3' "
         "desde la carpeta del juego original a la carpeta 'public/assets/' de la landing."),

        ("Fase 2: Motor de Audio y Sistema de Diseño",
         "Configuración de temas en Tailwind con las variables hexadecimales de Omnivis (#020713, #15355A, #00F5FF, #835F9E). "
         "Implementación del hook de audio global con soporte para autoplay amortiguado, persistencia de estado mute y visualizador gráfico."),

        ("Fase 3: Construcción de Componentes Principales",
         "Desarrollo secuencial de Navbar HUD interactivo, bloque Omnivis, Hero con mockup 3D tilt, "
         "grilla de características, galería táctica de sectores y el terminal interactivo embebido."),

        ("Fase 4: Micro-Animaciones, Efectos de Partículas y Sonido FX",
         "Incorporación de canvas con polvo estelar orbital, líneas de escaneo HUD, efectos magnéticos en botones "
         "y micro-beeps sintetizados al hacer click en elementos de interfaz."),

        ("Fase 5: Optimización para Vercel y Verificación",
         "Configuración de vercel.json, headers de compresión de audio e imágenes, prueba de compilación de producción ('npm run build') "
         "y comprobación de compatibilidad responsiva en resoluciones móviles, tablets y monitores ultra-wide.")
    ]

    for p_title, p_body in phases:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        r_t = p.add_run(p_title)
        r_t.bold = True
        r_t.font.size = Pt(10.5)
        r_t.font.color.rgb = COLOR_CYAN
        
        p_b = doc.add_paragraph(p_body)
        p_b.paragraph_format.left_indent = Inches(0.2)
        p_b.paragraph_format.space_after = Pt(6)

    # Final Notes
    doc.add_paragraph().paragraph_format.space_after = Pt(14)
    p_fin = doc.add_paragraph()
    p_fin.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_fin = p_fin.add_run("Omnivis © 2026 — Inteligencia que transforma tu entorno. Todos los derechos reservados.")
    r_fin.font.size = Pt(9)
    r_fin.font.italic = True
    r_fin.font.color.rgb = RGBColor(148, 163, 184)

    target_path = os.path.join(os.getcwd(), "PLAN_LANDING_PAGE_OMNIVIS.docx")
    doc.save(target_path)
    print(f"Document saved successfully at: {target_path}")

if __name__ == "__main__":
    create_document()
